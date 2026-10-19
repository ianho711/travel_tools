"""Real Chromium UI tests; requires Playwright and a locally installed browser.

Run: python tests/browser_test.py [--browser /usr/bin/chromium]
The test starts its own loopback HTTP server and writes artifacts outside the repo.
"""
import argparse
import functools
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
import json
from pathlib import Path
import shutil
import tempfile
import threading

from playwright.sync_api import sync_playwright, expect
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
STORE = 'travel-split-v1'


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *_args):
        pass


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--browser', default=shutil.which('chromium') or shutil.which('google-chrome'))
    args = parser.parse_args()
    artifacts = Path(tempfile.mkdtemp(prefix='travel-browser-'))
    handler = functools.partial(QuietHandler, directory=str(ROOT))
    server = ThreadingHTTPServer(('127.0.0.1', 0), handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    url = f'http://127.0.0.1:{server.server_port}/index.html'
    passed = []

    def record(message):
        passed.append(message)
        print('PASS:', message, flush=True)

    try:
        with sync_playwright() as pw:
            options = {'headless': True, 'args': ['--no-sandbox']}
            if args.browser:
                options['executable_path'] = args.browser
            browser = pw.chromium.launch(**options)
            context = browser.new_context(viewport={'width': 1440, 'height': 1000}, accept_downloads=True)
            page = context.new_page()
            errors, external_requests = [], []
            page.on('pageerror', lambda err: errors.append(str(err)))

            def route_request(route):
                if route.request.url.startswith(f'http://127.0.0.1:{server.server_port}/'):
                    route.continue_()
                else:
                    external_requests.append(route.request.url)
                    route.abort()

            context.route('**/*', route_request)
            page.goto(url)
            expect(page).to_have_title('旅伴帳本｜旅遊費用分攤')
            expect(page.get_by_role('tab')).to_have_count(12)
            expect(page.locator('#total-expense')).to_have_text('0.0元')
            expect(page.locator('#roster-summary')).to_contain_text('0 位已填姓名 / 30 個名額')
            record('初始 30 人／10 天，無假支出，頁面無外部依賴')

            def state():
                return json.loads(page.evaluate('(key) => localStorage.getItem(key)', STORE))

            def roster(opened=True):
                if page.locator('#roster').evaluate('(el) => el.open') != opened:
                    page.locator('#roster summary').click()

            def dialog(action, accept=True):
                messages = []

                def respond(d):
                    messages.append(d.message)
                    d.accept() if accept else d.dismiss()

                page.once('dialog', respond)
                with page.expect_event('dialog'):
                    action()
                assert messages, 'Expected a confirmation dialog'
                return messages[0]

            def tab(day_id):
                page.locator('#tab-' + day_id).click()

            def field(kind, eid):
                return page.locator(f'[data-field="{kind}"][data-expense="{eid}"]')

            def part(eid, pid):
                return page.locator(f'[data-field="participation"][data-expense="{eid}"][data-person="{pid}"]')

            def share(eid, pid, value):
                expect(page.locator(f'[data-share="{eid}-{pid}"]')).to_have_text(value)

            def all_button(eid, select=True):
                return page.locator(f'[data-action="{"select-all" if select else "select-none"}"][data-expense="{eid}"]')

            def set_fee(eid, amount, participants, payer):
                field('amount', eid).fill(amount)
                field('payer', eid).select_option(payer)
                for pid in participants:
                    part(eid, pid).check()

            def driver(day_id, pid):
                return page.locator(f'#driver-{day_id}-{pid}')

            roster()
            for i, name in enumerate(['阿安', '小晴', '阿哲']):
                page.locator('[data-field="member-name"]').nth(i).fill(name)
            roster(False)
            s = state()
            a, b, c = [m['id'] for m in s['members'][:3]]
            last = s['members'][29]['id']
            d1, d2 = s['days'][:2]
            e1 = {e['category']: e['id'] for e in d1['expenses']}
            e2 = {e['category']: e['id'] for e in d2['expenses']}
            tab(d1['id'])
            set_fee(e1['breakfast'], '100', [c, a, b], a)
            share(e1['breakfast'], a, '33.4')
            share(e1['breakfast'], b, '33.3')
            share(e1['breakfast'], c, '33.3')
            assert part(e1['breakfast'], a).evaluate('(el) => el.tagName === "INPUT" && el.type === "checkbox"')
            record('真實 checkbox 可操作，100 元／3 人顯示 33.4、33.3、33.3')

            set_fee(e1['lunch'], '90', [b, c], b)
            set_fee(e1['dinner'], '70', [a, c], c)
            share(e1['lunch'], a, '0.0')
            share(e1['lunch'], b, '45.0')
            share(e1['dinner'], a, '35.0')
            tab('overview')
            expect(page.locator('#total-expense')).to_have_text('260.0元')
            expect(page.locator(f'[data-summary="{a}-meals"]')).to_have_text('68.4')
            expect(page.locator('#total-difference')).to_have_text('0.0元')
            record('三餐勾選互相獨立，總覽立即反映餐費與代墊')

            with page.expect_download() as info:
                page.get_by_role('button',name='匯出試算表',exact=True).click()
            report_file=artifacts/'meal-report.xlsx'
            info.value.save_as(report_file)
            report=load_workbook(report_file,data_only=False)
            assert report['整趟總覽']['B3'].value==260
            assert report['整趟總覽']['A7'].value=='阿安'
            assert report['整趟總覽']['B7'].value==68.4
            assert report['整趟總覽']['M7'].value==31.6
            assert report['每餐參與名單']['E4'].value=='阿安、小晴、阿哲'
            assert report['每餐參與名單']['E5'].value=='小晴、阿哲'
            assert report['每餐參與名單']['E6'].value=='阿安、阿哲'
            daily_report=report['第 1 天']
            assert daily_report['C13'].value=='參與' and daily_report['D13'].value==33.4
            assert daily_report['E13'].value=='未參與' and daily_report['F13'].value==0
            assert daily_report['G13'].value=='參與' and daily_report['H13'].value==35
            assert not any(c.data_type=='f' for sh in report for row in sh for c in row)
            record('實際匯出 XLSX：三餐參與名單、沒吃零元、尾差及總覽正確，報表不需程式或公式')

            tab(d1['id'])
            page.locator('[data-field="use-car"]').check()
            driver(d1['id'], a).check()
            set_fee(e1['fuel'], '100', [a, b, c], a)
            set_fee(e1['rental'], '120', [a, b, c], a)
            share(e1['fuel'], a, '0.0')
            share(e1['rental'], a, '0.0')
            share(e1['fuel'], b, '50.0')
            share(e1['breakfast'], a, '33.4')
            driver(d1['id'], a).uncheck()
            expect(page.locator('#expense-' + e1['fuel'])).to_contain_text('尚未設定具名司機')
            driver(d1['id'], b).check()
            share(e1['fuel'], a, '50.0')
            share(e1['fuel'], b, '0.0')
            driver(d1['id'], a).check()
            share(e1['fuel'], c, '100.0')
            driver(d1['id'], b).uncheck()
            record('變更司機、多位司機、缺司機提示、司機勾選車資免付且餐費照付')

            tab(d2['id'])
            page.locator('[data-field="use-car"]').check()
            driver(d2['id'], b).check()
            set_fee(e2['fuel'], '60', [a, b, c], b)
            share(e2['fuel'], a, '30.0')
            share(e2['fuel'], b, '0.0')
            tab(d1['id'])
            expect(driver(d1['id'], a)).to_be_checked()
            expect(driver(d1['id'], b)).not_to_be_checked()
            share(e1['fuel'], a, '0.0')
            tab(d2['id'])
            expect(driver(d2['id'], b)).to_be_checked()
            expect(part(e2['breakfast'], a)).not_to_be_checked()
            record('實際切換日期，A／B 每日司機與跨日用餐紀錄各自獨立')

            tab(d1['id'])
            dialog(lambda: page.locator('[data-field="use-car"]').click(), accept=False)
            expect(page.locator('[data-field="use-car"]')).to_be_checked()
            message = dialog(lambda: page.locator('[data-field="use-car"]').click())
            assert '仍會保留' in message
            expect(field('amount', e1['fuel'])).to_have_value('100')
            share(e1['fuel'], b, '0.0')
            expect(page.locator('#expense-' + e1['fuel'])).to_contain_text('當日未用車')
            tab('overview')
            expect(page.locator('#total-difference')).to_have_text('220.0元')
            tab(d1['id'])
            page.locator('[data-field="use-car"]').check()
            share(e1['fuel'], b, '50.0')
            record('取消用車須確認，可取消操作；費用不遺失、差額顯示，重新用車可恢復')

            field('amount', e1['breakfast']).fill('100.1')
            share(e1['breakfast'], a, '33.4')
            share(e1['breakfast'], b, '33.4')
            all_button(e1['lunch'], False).click()
            expect(page.locator('#expense-' + e1['lunch'])).to_contain_text('沒有可分攤者')
            share(e1['lunch'], b, '0.0')
            part(e1['lunch'], b).check()
            part(e1['lunch'], c).check()
            field('amount', e1['other']).fill('0')
            expect(page.locator('#expense-' + e1['other'] + ' .expense-status')).to_be_empty()
            field('amount', e1['other']).fill('-1')
            expect(field('amount', e1['other'])).to_have_attribute('aria-invalid', 'true')
            field('amount', e1['other']).fill('')
            record('修改金額即時計算，正金額無分攤者、0 元、空白與不合法金額正確處理')

            roster()
            page.locator('#name-' + last).fill('第30位旅伴')
            roster(False)
            all_button(e1['dinner'], False).click()
            part(e1['dinner'], last).check()
            field('payer', e1['dinner']).select_option(last)
            share(e1['dinner'], last, '70.0')
            record('第 30 位成員可真正點選、分攤及代墊')

            # Clearing a name must preserve its selected IDs; bulk operations touch named members only.
            roster()
            page.locator('#name-' + c).fill('   ')
            roster(False)
            expect(part(e1['breakfast'], c)).to_be_disabled()
            all_button(e1['breakfast'], False).click()
            expense = next(e for e in state()['days'][0]['expenses'] if e['id'] == e1['breakfast'])
            assert expense['participantIds'] == [c]
            all_button(e1['breakfast']).click()
            expect(page.locator(f'[data-count="{e1["breakfast"]}"]')).to_have_text('3')
            share(e1['breakfast'], c, '0.0')
            roster()
            page.locator('#name-' + c).fill('阿哲')
            roster(False)
            expect(page.locator(f'[data-count="{e1["breakfast"]}"]')).to_have_text('4')
            record('空白姓名不計人數，全選與取消全選只操作具名成員，原紀錄可恢復')

            tab('overview')
            before = page.locator(f'[data-summary^="{a}-"]').all_text_contents()
            roster()
            page.locator('#name-' + a).fill('小安改名')
            roster(False)
            assert page.locator(f'[data-summary^="{a}-"]').all_text_contents() == before
            tab(d1['id'])
            expect(driver(d1['id'], a)).to_be_checked()
            assert next(e for e in state()['days'][0]['expenses'] if e['id'] == e1['breakfast'])['payerId'] == a
            expect(part(e1['breakfast'], a)).to_be_checked()
            record('改姓名同步顯示，分攤、代墊、參與及司機 ID 不變')

            # Only the compact title sticks; the large editing form must scroll away.
            scroll = page.locator('#daily-scroll')
            scroll.evaluate('(el) => { el.scrollLeft=0; el.scrollTop=0; }')
            x = page.locator('#daily-scroll tbody .name-cell').first.bounding_box()['x']
            editor_height = page.locator('#daily-expense-settings').bounding_box()['height']
            scroll.evaluate('(el, top) => { el.scrollLeft=450; el.scrollTop=top; }', editor_height + 100)
            y = page.locator('#title-' + e1['breakfast']).bounding_box()['y']
            assert page.locator('#expense-' + e1['breakfast']).bounding_box()['y'] < scroll.bounding_box()['y'] - 100
            scroll.evaluate('(el) => { el.scrollTop+=200; }')
            assert abs(page.locator('#daily-scroll tbody .name-cell').first.bounding_box()['x'] - x) < 2
            assert abs(page.locator('#title-' + e1['breakfast']).bounding_box()['y'] - y) < 2
            assert page.locator('.daily-compact-row').bounding_box()['height'] <= 64
            # At a smaller laptop viewport, enough space remains for several participant rows.
            page.set_viewport_size({'width': 1280, 'height': 720})
            scroll.evaluate('(el, top) => { el.scrollLeft=0; el.scrollTop=top; el.scrollIntoView({block:"center"}); }', editor_height + 300)
            visible_rows = scroll.evaluate("""el => {
                const box=el.getBoundingClientRect(), head=el.querySelector('.daily-compact-row').getBoundingClientRect();
                return [...el.querySelectorAll('tbody tr')].filter(row=>{
                    const r=row.getBoundingClientRect();return r.top>=head.bottom&&r.bottom<=box.bottom;
                }).length;
            }""")
            assert visible_rows >= 4, visible_rows
            page.screenshot(path=str(artifacts / 'laptop-daily-scrolled.png'), full_page=True)
            record('費用編輯區會隨捲動移開，僅固定不足 64px 的精簡標題；720px 高度可完整看到至少四列人員')

            original = state()
            summary_before = page.evaluate('(s) => TripCore.summarize(s)', original)
            page.get_by_role('button', name='收合費用設定', exact=True).click()
            expect(page.locator('#daily-expense-settings')).to_be_hidden()
            assert state() == original
            assert page.locator('.daily-compact-row').bounding_box()['y'] - scroll.bounding_box()['y'] < 2
            part(e1['breakfast'], a).uncheck()
            part(e1['breakfast'], a).check()
            assert page.evaluate('(s) => TripCore.summarize(s)', state()) == summary_before
            original = state()
            page.screenshot(path=str(artifacts / 'laptop-daily-collapsed.png'), full_page=True)
            page.get_by_role('button', name='展開費用設定', exact=True).click()
            expect(page.locator('#daily-expense-settings')).to_be_visible()
            amount_before = field('amount', e1['breakfast']).input_value()
            field('amount', e1['breakfast']).fill('-1')
            page.get_by_role('button', name='收合費用設定', exact=True).click()
            tab('overview')
            page.locator(f'[data-action="goto-expense"][data-expense="{e1["breakfast"]}"]').click()
            expect(page.locator('#daily-expense-settings')).to_be_visible()
            expect(field('amount', e1['breakfast'])).to_be_focused()
            field('amount', e1['breakfast']).fill(amount_before)
            assert state() == original
            page.set_viewport_size({'width': 1440, 'height': 1000})
            scroll.evaluate('(el) => { el.scrollLeft=0; el.scrollTop=0; }')
            record('費用設定可展開／收合，收合後 checkbox 正常，資料不變；待處理入口會自動展開並聚焦金額')

            page.get_by_role('button', name='＋ 新增費用', exact=True).click()
            added = state()['days'][0]['expenses'][-1]['id']
            field('expense-name', added).fill('加點早餐')
            field('category', added).select_option('breakfast')
            set_fee(added, '30', [a, b], b)
            share(added, a, '15.0')
            expect(part(e1['breakfast'], last)).to_be_checked()
            dialog(lambda: page.locator(f'[data-action="delete-expense"][data-expense="{added}"]').click(), False)
            assert len(state()['days'][0]['expenses']) == 8
            dialog(lambda: page.locator(f'[data-action="delete-expense"][data-expense="{added}"]').click())
            assert len(state()['days'][0]['expenses']) == 7
            record('同類費用可新增多筆且分開勾選，有資料費用刪除可取消或確認')

            roster()
            page.get_by_role('button', name='＋ 新增成員', exact=True).click()
            added_person = state()['members'][-1]['id']
            page.locator('#name-' + added_person).fill('新增旅伴')
            delete_person = page.locator(f'[data-action="delete-member"][data-person="{added_person}"]')
            dialog(lambda: delete_person.click(), False)
            assert len(state()['members']) == 31
            dialog(lambda: delete_person.click())
            assert len(state()['members']) == 30
            roster(False)
            page.get_by_role('button', name='＋ 新增一天', exact=True).click()
            field_day = page.locator('[data-field="day-name"]')
            field_day.fill('額外一天')
            expect(page.get_by_role('tab', name='額外一天', exact=True)).to_have_attribute('aria-selected', 'true')
            dialog(lambda: page.get_by_role('button', name='刪除此天', exact=True).click(), False)
            assert len(state()['days']) == 11
            dialog(lambda: page.get_by_role('button', name='刪除此天', exact=True).click())
            assert len(state()['days']) == 10
            record('人數與天數可增減，日期名稱同步，刪除已有資料需確認')

            tab('overview')
            overview = page.locator('.overview-table').inner_text()
            before_export = state()
            with page.expect_download() as info:
                page.get_by_role('button', name='匯出唯讀網頁', exact=True).click()
            readonly_file = artifacts / 'trip-view.html'
            info.value.save_as(readonly_file)
            assert 'id="trip-snapshot"' in readonly_file.read_text()
            with page.expect_download() as info:
                page.get_by_role('button', name='匯出 JSON', exact=True).click()
            download = info.value
            backup = artifacts / 'backup.json'
            download.save_as(backup)
            assert json.loads(backup.read_text()) == before_export
            before_invalid = state()
            for payload in [b'{bad json', b'{"version":9}', json.dumps({**before_export, 'members': []}).encode()]:
                page.locator('#import-file').set_input_files({'name': 'invalid.json', 'mimeType': 'application/json', 'buffer': payload})
                expect(page.locator('#flash')).to_contain_text('匯入失敗，原資料未變更')
                assert state() == before_invalid
            record('真正下載 JSON 備份；不合法 JSON／版本／名單匯入失敗且不覆蓋原資料')

            tab(d1['id'])
            field('amount', e1['breakfast']).fill('200')
            modified = state()
            dialog(lambda: page.locator('#import-file').set_input_files(str(backup)), False)
            assert state() == modified
            dialog(lambda: page.locator('#import-file').set_input_files(str(backup)))
            expect(page.locator('#flash')).to_contain_text('已成功還原')
            assert state() == before_export
            assert page.locator('.overview-table').inner_text() == overview
            page.reload()
            assert state() == before_export
            assert page.locator('.overview-table').inner_text() == overview
            record('匯入可取消，成功還原後資料與總覽一致，重新整理仍保存')

            # Serve the exported file from memory, without placing any personal data in the repo.
            context.route('**/shared-view.html', lambda r: r.fulfill(path=str(readonly_file), content_type='text/html; charset=utf-8'))
            tab(d1['id'])
            field('amount', e1['breakfast']).fill('999')
            editor_state = state()
            shared = context.new_page()
            shared.goto(url.replace('index.html', 'shared-view.html'))
            expect(shared).to_have_title('旅伴帳本｜唯讀旅遊費用')
            expect(shared.locator('.tools')).to_be_hidden()
            expect(shared.locator('#roster')).to_be_hidden()
            expect(shared.locator('[data-action="add-day"]')).to_be_hidden()
            assert shared.locator('.overview-table').inner_text() == overview
            expect(shared.locator('#save-status')).to_contain_text('唯讀快照 · 更新於')
            shared.locator('#tab-' + d1['id']).click()
            expect(shared.locator('#amount-' + e1['breakfast'])).to_have_value('100.1')
            assert shared.locator('#view input:not(:disabled),#view select:not(:disabled)').count() == 0
            expect(shared.locator('#part-' + e1['breakfast'] + '-' + a)).to_be_checked()
            expect(shared.locator('[data-action="add-expense"]')).to_be_hidden()
            shared.reload()
            assert state() == editor_state, 'Readonly viewer must never touch editor localStorage'
            # Returning to overview uses only the embedded snapshot, not the newer editor state.
            shared.locator('#tab-overview').click()
            assert shared.locator('.overview-table').inner_text() == overview
            shared.close()
            dialog(lambda: page.locator('#import-file').set_input_files(str(backup)))
            record('唯讀 HTML 快照可獨立開啟及切換日期，不能操作編輯欄位，不讀寫管理者帳本')

            dialog(lambda: page.get_by_role('button', name='重設全部', exact=True).click(), False)
            assert state() == before_export
            dialog(lambda: page.get_by_role('button', name='載入範例', exact=True).click(), False)
            assert state() == before_export
            page.screenshot(path=str(artifacts / 'desktop-overview.png'), full_page=True)
            tab(d1['id'])
            page.screenshot(path=str(artifacts / 'desktop-daily.png'), full_page=True)
            page.set_viewport_size({'width': 390, 'height': 844})
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth + 1')
            part(e1['breakfast'], a).uncheck()
            expect(part(e1['breakfast'], a)).not_to_be_checked()
            part(e1['breakfast'], a).check()
            expect(part(e1['breakfast'], a)).to_be_checked()
            page.screenshot(path=str(artifacts / 'mobile-daily.png'), full_page=True)
            tab('overview')
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth + 1')
            page.screenshot(path=str(artifacts / 'mobile-overview.png'), full_page=True)
            record('重設／載入範例需確認，390px 手機寬度無頁面溢出且 checkbox 可操作')

            # Prove sample loading/reset work rather than merely showing dialogs.
            dialog(lambda: page.get_by_role('button', name='載入範例', exact=True).click())
            expect(page.locator('#issue-count')).to_have_text('0項')
            assert len(state()['members']) == 4 and len(state()['days']) == 2
            dialog(lambda: page.get_by_role('button', name='重設全部', exact=True).click())
            expect(page.locator('#total-expense')).to_have_text('0.0元')
            assert len(state()['members']) == 30 and len(state()['days']) == 10
            record('載入範例有有效帳目且無待處理，重設確實恢復空白 30 人／10 天')

            # New personal ledger: exercise real purpose, +/- controls and settlement transfers.
            roster()
            for i, name in enumerate(['甲旅伴', '乙旅伴', '丙旅伴']):
                page.locator('[data-field="member-name"]').nth(i).fill(name)
            roster(False)
            fresh = state()
            pa, pb, pc = [m['id'] for m in fresh['members'][:3]]
            first_day = fresh['days'][0]
            bill = first_day['expenses'][0]['id']
            tab(first_day['id'])
            set_fee(bill, '6740', [pa], pb)
            tab('personal')
            page.locator('#personal-person').select_option(pa)

            def add_entry(kind, description, amount, direction='plus', payer=None):
                page.get_by_role('button', name='＋ 新增個人紀錄', exact=True).click()
                entry_id = state()['ledger'][-1]['id']
                page.locator('#ledger-type-' + entry_id).select_option(kind)
                page.locator('#ledger-description-' + entry_id).fill(description)
                if kind == 'adjustment':
                    page.locator('#ledger-direction-' + entry_id).select_option(direction)
                page.locator('#ledger-amount-' + entry_id).fill(amount)
                if payer:
                    page.locator('#ledger-payer-' + entry_id).select_option(payer)
                return entry_id

            extra = add_entry('adjustment', '額外車資', '4000')
            expect(page.locator('[data-personal="pay"]')).to_have_text('10,740.0 元')
            page.locator('#ledger-direction-' + extra).select_option('minus')
            expect(page.locator('[data-personal="pay"]')).to_have_text('2,740.0 元')
            expect(page.locator('[data-personal="paid"]')).to_have_text('0.0 元')
            page.locator('#ledger-direction-' + extra).select_option('plus')
            page.locator('#ledger-amount-' + extra).fill('3540')
            expect(page.locator('[data-personal="pay"]')).to_have_text('10,280.0 元')
            dialog(lambda: page.locator(f'[data-action="delete-ledger"][data-ledger="{extra}"]').click(), False)
            assert len(state()['ledger']) == 1
            dialog(lambda: page.locator(f'[data-action="delete-ledger"][data-ledger="{extra}"]').click())
            expect(page.locator('[data-personal="pay"]')).to_have_text('6,740.0 元')
            record('個人用途、金額與正負可實際編輯，多筆紀錄與刪除確認正常，加減收不當作繳款')

            payment = add_entry('payment', '出發前先收', '10000')
            expect(page.locator('#ledger-direction-' + payment)).to_be_disabled()
            expect(page.locator('#ledger-direction-' + payment)).to_have_value('minus')
            expect(page.locator('[data-personal="receive"]')).to_have_text('3,260.0 元')
            refund = add_entry('refund', '已退回差額', '3260')
            expect(page.locator('[data-personal="receive"]')).to_have_text('0.0 元')
            expect(page.locator('[data-personal="pay"]')).to_have_text('0.0 元')
            tab('overview')
            expect(page.locator('#total-expense')).to_have_text('6,740.0元')
            expect(page.locator('#collected-total')).to_contain_text('6,740.0 元')
            tab('personal')
            charge = add_entry('charge', '個人加購', '350', payer=pb)
            expect(page.locator('[data-personal="pay"]')).to_have_text('350.0 元')
            page.locator('#personal-person').select_option(pb)
            expect(page.locator('[data-personal="advanced"]')).to_have_text('7,090.0 元')
            expect(page.locator('.advance-list')).to_contain_text('350.0 元')
            # A payer may be viewing an advance for someone else's purchase: follow the owner.
            page.locator(f'.advance-list [data-ledger="{charge}"]').click()
            expect(page.locator('#personal-person')).to_have_value(pa)
            expect(page.locator('#ledger-amount-' + charge)).to_be_focused()
            record('先收款、應退差額、實際退款與個人加購分開計算；代墊自動彙總且來源可點選')

            bad = add_entry('adjustment', '待核對', '-1')
            expect(page.locator('#ledger-amount-' + bad)).to_have_attribute('aria-invalid', 'true')
            tab('overview')
            expect(page.locator('#pending-issues')).to_contain_text('待核對')
            page.locator(f'[data-action="goto-expense"][data-ledger="{bad}"]').click()
            expect(page.locator('#ledger-amount-' + bad)).to_be_focused()
            page.locator('#ledger-amount-' + bad).fill('0')
            roster()
            page.locator('[data-field="member-name"]').first.fill('改名甲')
            roster(False)
            expect(page.locator('#personal-person')).to_have_value(pa)
            expect(page.locator('[data-personal="pay"]')).to_have_text('350.0 元')
            saved_personal = state()
            page.reload()
            expect(page.locator('#tab-personal')).to_have_attribute('aria-selected', 'true')
            page.locator('#personal-person').select_option(pa)
            assert state() == saved_personal
            expect(page.locator('[data-personal="pay"]')).to_have_text('350.0 元')
            record('個人非法金額有提示與修正入口；改名與重新整理不遺失紀錄')

            # Backup / export includes every personal row and source advance.
            with page.expect_download() as info:
                page.get_by_role('button', name='匯出 JSON', exact=True).click()
            personal_backup = artifacts / 'personal-backup.json'
            info.value.save_as(personal_backup)
            assert json.loads(personal_backup.read_text()) == saved_personal
            with page.expect_download() as info:
                page.get_by_role('button', name='匯出試算表', exact=True).click()
            personal_xlsx = artifacts / 'personal-report.xlsx'
            info.value.save_as(personal_xlsx)
            exported = load_workbook(personal_xlsx)
            assert exported['個人紀錄']['A4'].value == '改名甲'
            assert exported['個人紀錄']['G4'].value == -10000
            assert exported['個人紀錄']['G5'].value == 3260
            assert exported['整趟總覽']['J7'].value == 10000
            assert exported['整趟總覽']['K7'].value == 3260
            assert exported['整趟總覽']['L7'].value == 350
            assert exported['代墊明細']['C4'].value == 6740
            dialog(lambda: page.get_by_role('button', name='重設全部', exact=True).click())
            dialog(lambda: page.locator('#import-file').set_input_files(str(personal_backup)))
            assert state() == saved_personal
            record('實際匯出及還原含個人紀錄的 JSON；XLSX 含正負、收退款、總覽與代墊來源')

            tab('personal')
            page.locator('#personal-person').select_option(pa)
            with page.expect_download() as info:
                page.get_by_role('button', name='匯出唯讀網頁', exact=True).click()
            personal_view = artifacts / 'personal-view.html'
            info.value.save_as(personal_view)
            context.route('**/personal-view.html', lambda r: r.fulfill(path=str(personal_view), content_type='text/html; charset=utf-8'))
            viewer = context.new_page()
            viewer.goto(url.replace('index.html', 'personal-view.html'))
            viewer.locator('#tab-personal').click()
            expect(viewer.locator('#personal-person')).to_be_enabled()
            viewer.locator('#personal-person').select_option(pb)
            expect(viewer.locator('[data-personal="advanced"]')).to_have_text('7,090.0 元')
            assert viewer.locator('#view input:not(:disabled),#view select:not(:disabled):not(#personal-person)').count() == 0
            expect(viewer.locator('[data-action="add-ledger"]')).to_be_hidden()
            viewer.close()
            assert state() == saved_personal
            page.set_viewport_size({'width': 390, 'height': 844})
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth + 1')
            page.screenshot(path=str(artifacts / 'mobile-personal.png'), full_page=True)
            page.set_viewport_size({'width': 1440, 'height': 1000})
            page.screenshot(path=str(artifacts / 'desktop-personal.png'), full_page=True)
            record('唯讀快照可切換查看不同成員但不能修改，個人明細在手機寬度無整頁溢出')

            # Deletion must clean up both ownership and other people's payer references.
            roster()
            dialog(lambda: page.locator('[data-action="delete-member"]').nth(1).click(), False)
            assert state() == saved_personal
            dialog(lambda: page.locator('[data-action="delete-member"]').nth(1).click())
            assert all(e['payerId'] != pb and e['personId'] != pb for e in state()['ledger'])
            assert next(e for e in state()['ledger'] if e['id'] == charge)['payerId'] is None
            tab('overview')
            expect(page.locator('#pending-issues')).to_contain_text('代墊者')
            record('刪成員需確認，個人紀錄與代墊參照一併處理，不留未知 ID')

            assert not errors, errors
            assert not external_requests, external_requests
            record('整段操作無 JavaScript 例外，所有外部網路請求被攔截且實際為零')

            # A blocked storage environment must warn instead of claiming saved.
            unavailable = browser.new_context()
            unavailable.add_init_script("Storage.prototype.setItem = function() { throw new DOMException('Blocked', 'SecurityError'); };")
            p2 = unavailable.new_page()
            p2.goto(url)
            p2.locator('#roster summary').click()
            p2.locator('[data-field="member-name"]').first.fill('儲存測試')
            expect(p2.locator('#save-status')).to_have_text('儲存失敗，請立即匯出 JSON 備份')
            expect(p2.locator('#flash')).to_contain_text('修改尚在本頁')
            record('localStorage 被阻擋時不假稱已儲存，清楚提示匯出備份')
            unavailable.close()
            browser.close()
    finally:
        server.shutdown()
        server.server_close()
    report = {'browser': 'Chromium', 'passed': passed, 'artifacts': str(artifacts), 'limitations': ['file:// navigation is blocked by the environment browser policy; tested using loopback HTTP instead.', 'Firefox, Safari and real mobile devices were not tested.']}
    (artifacts / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(f'\n{len(passed)} browser scenarios passed. Artifacts: {artifacts}', flush=True)


if __name__ == '__main__':
    main()
