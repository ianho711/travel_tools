"""Build the uploadable XLSX from the same recipe used by Google Apps Script."""
import argparse
import html
import json
from pathlib import Path
import subprocess

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties

HERE = Path(__file__).resolve().parent


def export_setup_helper(output=None):
    output = Path(output or HERE/'啟用助手.html')
    before = '''<!doctype html>
<html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>旅伴帳本｜Google 試算表啟用助手</title>
<style>body{font:16px/1.8 system-ui,sans-serif;background:#f4f5ef;color:#233d37;margin:0;padding:24px}main{max-width:840px;margin:auto;background:white;padding:28px;border-radius:16px}h1{font-size:26px}li{margin:14px 0}button,.button{font:inherit;display:inline-block;padding:10px 18px;border:0;border-radius:8px;background:#176253;color:white;text-decoration:none;cursor:pointer}textarea{width:100%;height:190px;box-sizing:border-box;font:12px/1.5 monospace;border:1px solid #b9c9bd;border-radius:8px}#status{color:#176253;font-weight:700}.note{background:#fff7df;padding:14px;border-radius:8px}</style></head><body><main>
<h1>Google 試算表：你編輯，旅伴檢視</h1>
<p>這個助手只協助你複製設定程式，不會自動讀取或上傳 Google Drive 資料。</p>
<ol><li><a class="button" target="_blank" rel="noopener noreferrer" href="https://drive.google.com/drive/my-drive">開啟 Google Drive</a><br>在資料夾按「新增 → Google 試算表」，建立空白檔案。</li>
<li>在新試算表按「擴充功能 → Apps Script」。按下面的按鈕複製程式，取代編輯器內的預設內容並儲存。<br><button type="button" id="copy">複製設定程式</button><p id="status" role="status"></p></li>
<li>在函式下拉選單選 <strong>setupTravelSheet</strong>，按「執行」並完成 Google 授權。回到試算表重新整理，就會出現名單、10 天、總覽與原生核取方塊。</li>
<li>按試算表右上角「共用」，將旅伴設為<strong>檢視者</strong>，你保留編輯權限。以後修改會自動儲存，旅伴可看到最新資料。</li></ol>
<p class="note">若已上傳 XLSX，請先「另存為 Google 試算表」，再貼程式並執行同一函式。啟用會保留原有資料。這份工具不會替你更改資料夾的共用權限。</p>
<details><summary>設定程式（無法自動複製時，可手動全選複製）</summary><textarea id="code" readonly aria-label="Google Apps Script 設定程式">'''
    after = '''</textarea></details></main><script>
document.getElementById('copy').addEventListener('click',async()=>{
 const code=document.getElementById('code'),status=document.getElementById('status');
 try{await navigator.clipboard.writeText(code.value);status.textContent='已複製，請貼到 Apps Script 編輯器。';}
 catch{code.closest('details').open=true;code.focus();code.select();
  try{if(!document.execCommand('copy'))throw Error();status.textContent='已複製，請貼到 Apps Script 編輯器。';}
  catch{status.textContent='請在程式文字框按 Ctrl+A（Mac：⌘A），再複製。';}}
});
</script></body></html>'''
    output.write_text(before+html.escape((HERE/'Code.gs').read_text())+after)
    return output


def build_spec(model=None):
    command = ['node', str(HERE / 'build_spec.cjs')]
    if model:
        command.append(str(model))
    return json.loads(subprocess.run(command, check=True, capture_output=True, text=True).stdout)


def export_spec(spec, output):
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    wb.properties.title = '旅遊費用分攤｜Google 試算表範本'
    wb.properties.creator = '旅伴帳本'
    wb.defined_names.add(DefinedName('TravelMemberIDs', attr_text=f"'成員名單'!$A$5:$A${spec['n']+4}"))
    green = PatternFill('solid', fgColor='176253')
    pale = PatternFill('solid', fgColor='EDF4ED')
    yellow = PatternFill('solid', fgColor='FFF8D9')
    red = PatternFill('solid', fgColor='FCE8E6')
    edge = Side(style='thin', color='DBE4DB')
    for definition in spec['sheets']:
        ws = wb.create_sheet(definition['name'])
        for entry in definition['cells']:
            cell = ws.cell(entry['r'], entry['c'], entry['v'])
            if entry['kind'] == 'text':
                cell.data_type = 's'
            cell.font = Font(name='Noto Sans CJK TC', size=11, color='233D37')
            cell.alignment = Alignment(vertical='center', wrap_text=entry['r'] <= 11)
            if entry['kind'] == 'formula' or isinstance(entry['v'], (int, float)) and not isinstance(entry['v'], bool):
                cell.number_format = '#,##0.0'
        for a1 in definition['merges']:
            ws.merge_cells(a1)
        for col, width in definition['widths'].items():
            ws.column_dimensions[col].width = max(10, width / 7)
        for col in definition['hiddenColumns']:
            ws.column_dimensions[col].hidden = True
        for start, end in definition['hiddenRows']:
            for row in range(start, end + 1):
                ws.row_dimensions[row].hidden = True
        ws.freeze_panes = ws.cell(definition['freezeRows']+1, definition['freezeColumns']+1).coordinate
        header = 7 if ws.title == '整趟總覽' else 4 if ws.title == '成員名單' else 11 if ws.title.startswith('第 ') else 1
        for cell in ws[header]:
            cell.fill = green
            cell.font = Font(name='Noto Sans CJK TC', size=11, bold=True, color='FFFFFF')
        ws.row_dimensions[header].height = 30
        for a1 in definition['editable']:
            for row in ws[a1] if ':' in a1 else [(ws[a1],)]:
                for cell in row:
                    cell.fill = yellow
        for a1 in definition['checkboxes']:
            validation = DataValidation(type='custom', formula1=f'OR({a1.split(":")[0]}=TRUE,{a1.split(":")[0]}=FALSE)', allow_blank=True)
            validation.errorTitle = '請使用核取方塊'
            validation.error = '在 Google 試算表執行 setupTravelSheet，即可變成真正可點選的核取方塊。'
            validation.showErrorMessage = True
            ws.add_data_validation(validation)
            validation.add(a1)
        for entry in definition['validations']:
            if entry['type'] == 'list':
                validation = DataValidation(type='list', formula1='"'+','.join(entry['values'])+'"', allow_blank=False)
            elif entry['type'] == 'range':
                validation = DataValidation(type='list', formula1='TravelMemberIDs', allow_blank=True)
            else:
                validation = DataValidation(type='custom', formula1=entry['formula'], allow_blank=True)
            validation.showErrorMessage = True
            validation.errorTitle = '輸入格式錯誤'
            validation.error = '請使用下拉選項；金額須非負且最多一位小數。'
            ws.add_data_validation(validation)
            validation.add(entry['range'])
        if ws.title.startswith('第 '):
            for r in range(4, 10):
                ws.row_dimensions[r].height = 28
            ws.row_dimensions[10].height = 82
            for r in range(12, spec['last'] + 2):
                ws.row_dimensions[r].height = 27
                for cell in ws[r]:
                    cell.border = Border(bottom=edge)
                    if r % 2 == 0 and cell.fill.patternType is None:
                        cell.fill = pale
            for expense in range(len(spec['model']['days'][int(ws.title.split()[1])-1]['expenses'])):
                letter = get_column_letter(4 + 2 * expense)
                ws[f'{letter}10'].alignment = Alignment(wrap_text=True, vertical='top')
                ws[f'{letter}10'].font = Font(name='Noto Sans CJK TC', size=10, color='A82F33')
        if ws.title == '整趟總覽':
            for r in [1, 2, 3, 4, 5]:
                ws.row_dimensions[r].height = 30
            for r in range(8, spec['n'] + 9):
                ws.row_dimensions[r].height = 28
                for cell in ws[r]:
                    cell.border = Border(bottom=edge)
                    if r % 2 == 0:
                        cell.fill = pale
            for row in range(spec['n']+13, ws.max_row+1):
                ws.row_dimensions[row].height = 48
                ws.cell(row, 5).alignment = Alignment(wrap_text=True, vertical='center')
                ws.cell(row, 5).font = Font(name='Noto Sans CJK TC', size=10, color='A82F33')
            for cell in ['C3', 'F3', 'I3', 'C4', 'F4']:
                ws[cell].font = Font(name='Noto Sans CJK TC', size=14, bold=True, color='176253')
            ws.conditional_formatting.add('I3', CellIsRule(operator='greaterThan', formula=['0'], fill=red))
        if ws.title == '使用說明':
            for r in range(1, ws.max_row+1):
                ws.row_dimensions[r].height = 42
                ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='center')
        if ws.title == '成員名單':
            for r in range(5, spec['n']+5):
                ws.row_dimensions[r].height = 28
        if definition['hidden']:
            ws.sheet_state = 'hidden'
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.active = 0
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--model', type=Path)
    parser.add_argument('--output', type=Path, default=HERE/'旅遊費用分攤_Google試算表範本.xlsx')
    args = parser.parse_args()
    print(export_spec(build_spec(args.model), args.output))
    print(export_setup_helper())
