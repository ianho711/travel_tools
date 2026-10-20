"""Upgrade a private v1.9.0 ZIP to the current app version without opening HTML or driving a browser.

Usage: python scripts/repair_v190_package.py ORIGINAL.zip NEW.zip
Requires Node.js and openpyxl. Original inputs are never modified. Do not commit
private inputs/outputs to Git. Both HTMLs receive the same current scripts.
"""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path, PurePosixPath
import re
import subprocess
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
NODE=r'''
const fs=require('node:fs'),vm=require('node:vm'),{TextEncoder}=require('node:util'),assert=require('node:assert/strict');
const [root,input,output,stamp]=process.argv.slice(1),html=fs.readFileSync(root+'/index.html','utf8');
const extract=(s,id)=>s.match(new RegExp('<script id="'+id+'">([\\s\\S]*?)</script>'))[1];
const context=vm.createContext({TextEncoder}),C=vm.runInContext(extract(html,'trip-core')+';TripCore',context),E=vm.runInContext(extract(html,'trip-export')+';TripExport',context);
const originalRaw=fs.readFileSync(input+'/book.json','utf8'),state=C.decode(originalRaw),now=new Date(stamp);
const originalHTML=fs.readFileSync(input+'/editable.html','utf8'),old=vm.runInNewContext(extract(originalHTML,'trip-core')+';TripCore');
const before=old.summarize(old.decode(originalRaw)),after=C.summarize(state);
function existingFields(value,shape){if(Array.isArray(shape)){assert.equal(value.length,shape.length);return shape.map((item,i)=>existingFields(value[i],item));}if(shape&&typeof shape==='object')return Object.fromEntries(Object.keys(shape).map(key=>[key,existingFields(value[key],shape[key])]));return value;}
assert.equal(JSON.stringify(existingFields(after,before)),JSON.stringify(before),'Original financial results changed');
assert.equal(JSON.stringify(C.decode(C.encode(state))),JSON.stringify(state),'JSON round trip');
const originalView=fs.readFileSync(input+'/readonly.html','utf8'),snapshot=JSON.parse(originalView.match(/<script id="trip-snapshot"[^>]*>([\s\S]*?)<\/script>/)[1]);
assert.equal(JSON.stringify(C.validateState(snapshot.trip)),JSON.stringify(state),'Original readonly account differs from JSON');
// Execute initialization in a Node VM with a tiny document stub. No page is opened,
// no layout/render function is run, and no UI events are dispatched.
function bootstrap(page,stored,readonly) {
 const nodes=new Map(),make=()=>({textContent:'',addEventListener(){}});
 for(const id of ['trip-snapshot','trip-initial']){
  const match=page.match(new RegExp('<script id="'+id+'"[^>]*>([\\s\\S]*?)</script>'));
  if(match)nodes.set(id,{...make(),textContent:match[1]});
 }
 const data=new Map(stored===null?[]:[['travel-split-v1',stored]]);
 const store={getItem(key){if(readonly)throw Error('Readonly touched localStorage');return data.get(key)??null;},setItem(key,value){if(readonly)throw Error('Readonly wrote localStorage');data.set(key,value);}};
 const document={title:'',getElementById(id){if(['trip-snapshot','trip-initial'].includes(id))return nodes.get(id)||null;if(!nodes.has(id))nodes.set(id,make());return nodes.get(id);},querySelector:()=>make(),addEventListener(){}};
 const ctx=vm.createContext({document,localStorage:store,addEventListener(){},navigator:{},TextEncoder});
 vm.runInContext(extract(page,'trip-core')+'\n'+extract(page,'trip-storage'),ctx);
 const app=extract(page,'trip-app').replace('  render();if(initialMessage)','  globalThis.loadedState=state;globalThis.loadedReadOnly=readOnly;if(initialMessage)');
 vm.runInContext(app,ctx);assert.equal(ctx.loadedReadOnly,readonly);return JSON.stringify(ctx.loadedState);
}
// Update shared logic, version labels, the recovery banner and embedded account seed.
for(const [source,target,readonly] of [[originalHTML,'editable.html',false],[originalView,'readonly.html',true]]){
 let page=source;
 page=page.replace(/<style>[\s\S]*?<\/style>/,()=>html.match(/<style>[\s\S]*?<\/style>/)[0]);
 assert.equal(page.match(/<style>[\s\S]*?<\/style>/)[0],html.match(/<style>[\s\S]*?<\/style>/)[0],'Shared style mismatch');
 page=page.replace(/(<title>旅伴帳本 v)[^｜<]+/,()=>'<title>旅伴帳本 v'+C.APP_VERSION);
 page=page.replace(/(<span[^>]*id="app-version"[^>]*>)[^<]*(<\/span>)/,(_all,start,end)=>start+'v'+C.APP_VERSION+end);
 assert.ok(page.includes('<title>旅伴帳本 v'+C.APP_VERSION+'｜'),'Static title version mismatch');
 assert.ok(page.match(/id="app-version"[^>]*>v([0-9.]+)</)[1]===C.APP_VERSION,'Visible version mismatch');
 for(const id of ['trip-core','trip-export','trip-app'])page=page.replace(new RegExp('<script id="'+id+'">[\\s\\S]*?</script>'),()=>'<script id="'+id+'">'+extract(html,id)+'</script>');
 const storage='<script id="trip-storage">'+extract(html,'trip-storage')+'</script>';
 page=page.replace(/<script id="trip-storage">[\s\S]*?<\/script>\s*/,'').replace('<script id="trip-app">',()=>storage+'\n\n<script id="trip-app">');
 page=page.replace(/\s*<div id="storage-conflict"[^\n]*\n/,'\n');
 const banner=html.match(/  <div id="storage-conflict"[^\n]*\n/)[0];
 assert.ok(page.includes('<main>'),'Missing main element');page=page.replace('<main>',()=>banner+'  <main>');
 page=page.replace(/<script id="trip-initial"[^>]*>[\s\S]*?<\/script>\s*/,'');
 const json=JSON.stringify({createdAt:stamp,trip:state}).replace(/</g,'\\u003c');
 if(readonly)page=page.replace(/<script id="trip-snapshot"[^>]*>[\s\S]*?<\/script>/,()=>'<script id="trip-snapshot" type="application/json">'+json+'</script>');
 else page=page.replace('<script id="trip-app">',()=>'<script id="trip-initial" type="application/json">'+json+'</script>\n<script id="trip-app">');
 for(const id of ['trip-core','trip-export','trip-storage','trip-app']){
  assert.equal(extract(page,id),extract(html,id),'Shared code mismatch: '+id);new vm.Script(extract(page,id));
 }
 const embedded=JSON.parse(page.match(new RegExp('<script id="'+(readonly?'trip-snapshot':'trip-initial')+'"[^>]*>([\\s\\S]*?)</script>'))[1]);
 const restored=C.validateState(embedded.trip);assert.equal(JSON.stringify(restored),JSON.stringify(state));
 assert.equal(bootstrap(page,null,readonly),JSON.stringify(state),'Actual initialization differs from account');
 if(!readonly){const existing=JSON.parse(C.encode(state));existing.members[0].name='Local existing account';assert.equal(bootstrap(page,JSON.stringify(existing),false),JSON.stringify(existing),'Existing localStorage must take priority');}

 assert.equal(JSON.stringify(E.report(restored,now)),JSON.stringify(E.report(state,now)),'Editable/readonly report mismatch');
 fs.writeFileSync(output+'/'+target,page);
}
fs.writeFileSync(output+'/report.xlsx',Buffer.from(E.xlsx(state,now)));
fs.writeFileSync(output+'/tables.json',JSON.stringify(E.report(state,now)));
fs.writeFileSync(output+'/summary.json',JSON.stringify({appVersion:C.APP_VERSION,dataVersion:state.version,cashExpense:after.cashExpense/10,income:after.income/10,total:after.total/10,allocated:after.allocated/10,difference:after.difference/10,issues:after.issues.length,transportAllocated:after.transport.allocated/10,carAllocated:after.expenses.filter(e=>e.dayId==='vehicles').reduce((n,e)=>n+e.allocated,0)/10}));
'''


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('original',type=Path)
    parser.add_argument('output',type=Path)
    args=parser.parse_args()
    src=args.original.resolve();dest=args.output.resolve()
    if src==dest or dest.exists():raise ValueError('Output must be a new ZIP; never overwrite original or existing delivery')
    digest=hashlib.sha256(src.read_bytes()).hexdigest()
    with ZipFile(src) as z:
        if z.testzip() is not None:raise ValueError('Corrupt original ZIP')
        contents={i.filename:z.read(i) for i in z.infolist() if not i.is_dir()}
    for name in contents:
        path=PurePosixPath(name)
        if path.is_absolute() or '..' in path.parts or '\\' in name:raise ValueError('Unsafe ZIP path')
    htmls=[n for n in contents if n.endswith('.html')]
    views=[n for n in htmls if b'id="trip-snapshot"' in contents[n]]
    editable=[n for n in htmls if n not in views]
    books=[n for n in contents if n.endswith('.json')]
    excels=[n for n in contents if n.endswith('.xlsx')]
    if not(len(editable)==len(views)==len(books)==len(excels)==1):raise ValueError('Expected one editable HTML, readonly HTML, JSON and XLSX')
    with tempfile.TemporaryDirectory(prefix='travel-private-repair-') as tmp:
        tmp=Path(tmp);original=tmp/'original';output=tmp/'fixed';original.mkdir();output.mkdir()
        for target,source in [('editable.html',editable[0]),('readonly.html',views[0]),('book.json',books[0])]:
            (original/target).write_bytes(contents[source])
        stamp=datetime.now(timezone.utc).isoformat()
        subprocess.run(['node','-e',NODE,str(ROOT),str(original),str(output),stamp],check=True)
        wb=load_workbook(output/'report.xlsx',data_only=False)
        tables=json.loads((output/'tables.json').read_text());count=0
        assert wb.sheetnames==[t['name'] for t in tables]
        for table in tables:
            sh=wb[table['name']]
            for r,row in enumerate(table['rows'],1):
                for c,entry in enumerate(row,1):
                    actual=sh.cell(r,c).value;expected=entry['value']
                    assert actual==expected or expected=='' and actual is None,(sh.title,r,c)
                    if expected!='':count+=1
        assert not any(c.data_type=='f' for sh in wb for row in sh for c in row)
        summary=json.loads((output/'summary.json').read_text())
        def sum_column(name,col):
            return round(sum(row[col-1] for row in wb[name].iter_rows(min_row=4,values_only=True) if isinstance(row[col-1],(int,float))),1)
        car=summary.pop('carAllocated');transport=summary.pop('transportAllocated')
        assert sum_column('車資逐日核對',7)==sum_column('車資每日小計',4)==sum_column('車資個人小計',6)==car
        assert sum_column('每日交通費',7)==sum_column('每日交通費小計',6)==round(transport-car,1)
        assert sum_column('各車交通費',4)==wb['交通總計']['L4'].value
        assert wb['交通總計']['E4'].value==transport
        assert wb['整趟總覽']['J3'].value==summary['cashExpense'] and wb['整趟總覽']['L3'].value==summary['income']
        assert wb['整趟總覽']['B3'].value==wb['整趟總覽']['D3'].value==summary['total']==summary['allocated']
        assert wb['整趟總覽']['F3'].value==summary['difference']
        contents[editable[0]]=(output/'editable.html').read_bytes()
        contents[views[0]]=(output/'readonly.html').read_bytes()
        contents[excels[0]]=(output/'report.xlsx').read_bytes()
        summary.update({'originalZipSHA256':digest,'accountJSONUnchanged':True,'sharedScriptsIdentical':True,'editableReadonlyExcelMatch':True,'initializationVerifiedWithoutUI':True,'sheets':len(wb.sheetnames),'nonemptyCellsChecked':count,'verification':'Nonvisual Node VM, JSON round trip, openpyxl cell/column checks; no browser/UI/screenshots'})
        parent=PurePosixPath(books[0]).parent
        contents[str(parent/'修正版驗證結果.json')]=json.dumps(summary,ensure_ascii=False,indent=2).encode()
        manual=f'''旅伴帳本 v{summary['appVersion']} 修正版（2026-10-20）
JSON資料格式：version {summary['dataVersion']}，保留原備份相容性。

本版新增：按組贊助金、緊湊連續總覽。另保留多分頁衝突保護、車輛100筆上限與交通報表明細／小計分離。
贊助者增加應負擔，勾選的其他家庭每列算1組抵減，不增加支出、收入或代墊。真正交款後另記已繳款。
分攤規則與原JSON內容均不變；可編輯版、唯讀快照及Excel含相同帳目。
原ZIP未修改。此包含私人帳目，請自行選擇分享對象，不要上傳公開GitHub。

開始：關閉舊版編輯分頁，再解壓新包。可編輯HTML在該儲存空間沒有帳本時，載入隨附帳目；已有localStorage帳本時優先保留原帳，必要時先匯出備份再匯入隨附JSON。唯讀版只讀快照。
自動儲存需要支援Web Locks的新版瀏覽器；不支援時會保留草稿、停止自動寫入並提示匯出JSON。若本機檔案受瀏覽器限制，可在解壓目錄執行 python -m http.server 8000 --bind 127.0.0.1 ，再從同一網址開兩個分頁測試。

供使用者執行的手動確認（本次未操作UI、開啟網頁或截圖）：
1. 先備份。在相同檔案／網址開A、B兩頁；A改費用並等候儲存成功，B再改。B應顯示衝突、保留本頁修改且不覆蓋A。B完成視窗儲存到本頁後匯出JSON並確認下載，再按重新載入；若要保留B修改，需比對後手動合併，沒有自動合併。
2. 車輛超過100筆時應保留視窗草稿，顯示目前筆數及需新增筆數；一般編輯車輛後可正常匯出及還原JSON。
3. Excel：車資逐日核對G欄、車資每日小計D欄、車資個人小計F欄分別加總應相等。每日交通費G欄與每日交通費小計F欄應相等。各車交通費D欄加總等於交通總計L4。每張表分開核對，各表是相同費用的不同層級，不能把不同表的總額再次相加。
4. 此包帳目來自先前上傳的原始ZIP；若瀏覽器已有較新的修改，先匯出目前JSON，再於新版匯入，不要用舊包帳目取代最新帳。比較可編輯版、唯讀版、Excel總覽；原帳實際支出{summary['cashExpense']:g}、優惠收入{summary['income']:g}、分攤淨額及已分攤{summary['total']:g}、差額{summary['difference']:g}元。既有瀏覽器若是不同帳本，先備份，再匯入隨附JSON核對。
5. 到整趟總覽新增贊助金，選贊助者、填金額，確認自己的家庭禁選、其餘家庭可勾選；儲存後確認贊助者增加負擔、受惠家庭等額抵減。按取消不應更改帳目；交款未記錄前不應增加已繳款。
6. HTML及Excel版面、按鈕與提示由你確認；自動驗證結果見修正版驗證結果.json。
'''
        contents[str(parent/'修正版說明與手動確認.txt')]=manual.encode('utf-8-sig')
        # Rename only archive paths; the original JSON bytes remain untouched.
        renamed=lambda name:re.sub(r'(?<![0-9])1\.9\.0(?![0-9])',summary['appVersion'],name)
        renamed_contents={renamed(name):data for name,data in contents.items()}
        assert len(renamed_contents)==len(contents),'Archive path collision'
        contents=renamed_contents
        dest.parent.mkdir(parents=True,exist_ok=True)
        with ZipFile(dest,'x',ZIP_DEFLATED) as z:
            for name,data in contents.items():z.writestr(name,data)
        with ZipFile(dest) as z:
            assert z.testzip() is None
            assert z.read(renamed(books[0]))==(original/'book.json').read_bytes()
        assert hashlib.sha256(src.read_bytes()).hexdigest()==digest,'Original ZIP changed'
        print(json.dumps(summary,ensure_ascii=False,indent=2));print('Output:',dest)


if __name__=='__main__':main()
